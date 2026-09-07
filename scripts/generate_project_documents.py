from __future__ import annotations

import subprocess
from pathlib import Path
from datetime import date
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = DOCS / "document-assets"
ASSETS.mkdir(exist_ok=True)

NAVY = "17365D"; BLUE = "2F75B5"; CYAN = "DDEBF7"; PALE = "F3F6FA"
GREEN = "70AD47"; AMBER = "FFC000"; GREY = "666666"; WHITE = "FFFFFF"

def font(size=24, bold=False):
    candidates = [Path("C:/Windows/Fonts/aptos.ttf"), Path("C:/Windows/Fonts/calibri.ttf")]
    if bold:
        candidates = [Path("C:/Windows/Fonts/aptos-bold.ttf"), Path("C:/Windows/Fonts/calibrib.ttf")] + candidates
    for p in candidates:
        if p.exists(): return ImageFont.truetype(str(p), size)
    return ImageFont.load_default()

def box(draw, xy, text, fill=CYAN, outline=BLUE, size=22):
    draw.rounded_rectangle(xy, 15, fill="#"+fill, outline="#"+outline, width=3)
    x1,y1,x2,y2=xy
    lines=text.split("\n")
    total=sum(draw.textbbox((0,0),s,font=font(size,True))[3] for s in lines)+8*(len(lines)-1)
    y=(y1+y2-total)/2
    for s in lines:
        f=font(size,True); bb=draw.textbbox((0,0),s,font=f)
        draw.text(((x1+x2-(bb[2]-bb[0]))/2,y),s,font=f,fill="#17365D")
        y += bb[3]-bb[1]+8

def arrow(draw, a, b, color=BLUE):
    draw.line([a,b], fill="#"+color, width=5)
    import math
    ang=math.atan2(b[1]-a[1],b[0]-a[0]); l=16
    pts=[b,(b[0]-l*math.cos(ang-.5),b[1]-l*math.sin(ang-.5)),(b[0]-l*math.cos(ang+.5),b[1]-l*math.sin(ang+.5))]
    draw.polygon(pts, fill="#"+color)

def diagram(name, title, blocks, arrows, size=(1800,850), bands=None):
    im=Image.new("RGB",size,"white"); d=ImageDraw.Draw(im)
    d.text((50,25),title,font=font(34,True),fill="#17365D")
    if bands:
        for xy,label,color in bands:
            d.rounded_rectangle(xy,18,fill="#"+color,outline="#A6A6A6",width=2)
            d.text((xy[0]+15,xy[1]+10),label,font=font(20,True),fill="#555555")
    for xy,text_,fill in blocks: box(d,xy,text_,fill)
    for a,b in arrows: arrow(d,a,b)
    p=ASSETS/name; im.save(p,dpi=(160,160)); return p

def diagrams():
    user=diagram("architecture-user-workflow.png","User journey — clear gates, safe fallbacks",
      [((55,180,300,330),"Enroll\n(consent)","E2F0D9"),((360,180,605,330),"Identify\nor Unknown",CYAN),((665,180,910,330),"Calibrate\neyes + pose",CYAN),((970,180,1215,330),"Monitor\nquality + state",CYAN),((1275,180,1520,330),"Alert\npolicy", "FFF2CC"),((1275,500,1520,650),"Statistics\n& review","E2F0D9")],
      [((300,255),(360,255)),((605,255),(665,255)),((910,255),(970,255)),((1215,255),(1275,255)),((1395,330),(1395,500)),((1275,575),(1100,575)),((1100,575),(1100,330))])
    runtime=diagram("architecture-runtime.png","Runtime architecture — provider detail stops at semantic contracts",
      [((35,220,250,370),"Camera / video\nsource",CYAN),((300,220,515,370),"Capture +\nsource clock",CYAN),((565,220,780,370),"Providers\nYuNet / MP / PFLD","FFF2CC"),((830,220,1060,370),"Semantic\nobservations","E2F0D9"),((1110,220,1335,370),"Metrics +\ntemporal FSMs","E2F0D9"),((1385,110,1740,250),"Display / focus AOI",PALE),((1385,300,1740,440),"Trace / statistics",PALE),((1385,490,1740,630),"Alert adapter (planned)","FCE4D6")],
      [((250,295),(300,295)),((515,295),(565,295)),((780,295),(830,295)),((1060,295),(1110,295)),((1335,280),(1385,180)),((1335,310),(1385,370)),((1335,335),(1385,560))])
    portable=diagram("architecture-portability.png","Portability boundary — shared C++ core, thin target adapters",
      [((100,180,760,650),"PLATFORM-INDEPENDENT C++17 CORE\n\nSemantic observations\nQuality gates + calibration\nScheduler + latest-frame slot\nMetrics + FSMs\nConfiguration + statistics","E2F0D9"),((930,155,1650,300),"Windows x64\nOpenCV camera • LoadLibrary • CNG/DPAPI",CYAN),((930,360,1650,505),"Linux / Orin aarch64\nV4L2 • dlopen • OpenSSL • device telemetry",CYAN),((930,565,1650,710),"TI AM62 / Pi 5 (planned)\nBSP • camera • acceleration • packaging","FFF2CC")],
      [((760,270),(930,225)),((760,410),(930,430)),((760,555),(930,635))])
    tech=diagram("architecture-technology.png","Technology stack and release evidence",
      [((60,170,390,330),"Application\nC++17 • CMake",CYAN),((470,170,800,330),"Vision\nOpenCV 4.8",CYAN),((880,170,1210,330),"Optional landmarks\nMediaPipe 0.10.33","FFF2CC"),((1290,170,1620,330),"Models\nchecksum-pinned","FFF2CC"),((260,500,620,660),"Target package\nDLL / SO manifest",PALE),((720,500,1080,660),"SBOM + notices\nprovenance gate","FCE4D6"),((1180,500,1540,660),"Validation evidence\nWindows / Orin", "E2F0D9")],
      [((390,250),(470,250)),((800,250),(880,250)),((1210,250),(1290,250)),((635,330),(440,500)),((1045,330),(900,500)),((1455,330),(1360,500)),((620,580),(720,580)),((1080,580),(1180,580))])
    tree=diagram("developer-map.png","Repository map — where a new feature usually lands",
      [((60,150,390,300),"include/\npublic contracts",CYAN),((470,150,800,300),"src/\nimplementations + apps",CYAN),((880,150,1210,300),"tests/\ndeterministic gates","E2F0D9"),((1290,150,1620,300),"docs/\nclaims + evidence",PALE),((260,500,620,650),"mediapipe/\noptional C ABI bridge","FFF2CC"),((720,500,1080,650),"scripts/\nbuild, package, score",PALE),((1180,500,1540,650),"config/ + models/\nruntime inputs","FFF2CC")],
      [((390,225),(470,225)),((800,225),(880,225)),((1210,225),(1290,225)),((635,300),(440,500)),((965,300),(900,500)),((1375,300),(1360,500))])
    return user,runtime,portable,tech,tree

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def set_repeat_header(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),"true"); trPr.append(el)

def setup(doc, title, subtitle, landscape=False):
    sec=doc.sections[0]
    if landscape:
        sec.orientation=WD_ORIENT.LANDSCAPE; sec.page_width=Inches(11.69); sec.page_height=Inches(8.27)
    else:
        sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.62); sec.left_margin=sec.right_margin=Inches(.65)
    styles=doc.styles
    styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9.5); styles['Normal'].font.color.rgb=RGBColor.from_string('333333')
    for n,size,color in [('Title',30,NAVY),('Heading 1',19,NAVY),('Heading 2',14,BLUE),('Heading 3',11,NAVY)]:
        styles[n].font.name='Aptos Display'; styles[n].font.size=Pt(size); styles[n].font.color.rgb=RGBColor.from_string(color)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.space_after=Pt(5)
    r=p.add_run(title); r.bold=True; r.font.size=Pt(30); r.font.color.rgb=RGBColor.from_string(NAVY)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(subtitle); r.font.size=Pt(14); r.font.color.rgb=RGBColor.from_string(BLUE)
    doc.add_paragraph()
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"Engineering checkpoint • {date.today().isoformat()}\n").bold=True
    p.add_run("DMS Next • Sponsor and development use")
    doc.add_paragraph()
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("DMS Next  •  ")
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

def add_table(doc, headers, rows, widths=None, font_size=8):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,NAVY); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.color.rgb=RGBColor(255,255,255); r.bold=True; r.font.size=Pt(font_size)
    set_repeat_header(t.rows[0])
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=str(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            if ri%2: shade(cells[i],PALE)
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(1)
                for r in p.runs: r.font.size=Pt(font_size)
        if widths:
            for c,w in zip(cells,widths): c.width=Inches(w)
    return t

def add_picture(doc,p,width=9.5):
    par=doc.add_paragraph(); par.alignment=WD_ALIGN_PARAGRAPH.CENTER; par.add_run().add_picture(str(p),width=Inches(width))

def git_files():
    out=subprocess.check_output(['git','-c',f'safe.directory={ROOT.as_posix()}','ls-files'],cwd=ROOT,text=True)
    return [Path(x) for x in out.splitlines()]

PURPOSE={
 'AppPaths':'resolves executable-relative model and package paths', 'BackendEyeMapper':'normalizes provider landmarks into semantic eye contours',
 'BackendFaceGeometryMapper':'maps provider output to provider-neutral facial geometry', 'BackendGazeMapper':'derives semantic gaze inputs',
 'BackendOptions':'parses backend/provider command-line selection', 'BenchmarkOptions':'parses benchmark, input, ROI and profiling options',
 'BlinkTracker':'legacy EAR blink state tracker used by the camera/demo path', 'DisplayAoiAdapter':'creates display-only full/face/eyes/mouth views',
 'DmsEyeMetrics':'quality-gated per-driver eye calibration and eye temporal metrics', 'DmsEyeQualityAssessor':'scores eye ROI visibility/usability',
 'DmsEyeStatistics':'cumulative and configurable rolling eye-open/blink statistics', 'DmsHeadPoseEstimator':'head-pose estimation and neutral-pose calibration',
 'DmsObservation':'timestamped semantic observation and freshness contracts', 'DmsPolicy':'named operational policy defaults',
 'DmsPresentationConfig':'typed presentation, processing ROI and statistics configuration', 'DmsPresentationConfigLoader':'strict schema-versioned configuration loader',
 'DmsScheduler':'independent task cadence and depth-one latest-frame scheduling', 'DmsTemporalEvents':'yawn, pose, gaze, presence, availability and drowsiness FSMs',
 'DriverIdentity':'provider-neutral identity, embedding, PAD and store interfaces', 'DriverIdentityMatcher':'open-set candidate matching and ambiguity handling',
 'DriverProfileDatabase':'bounded profile serialization, import conflicts and deletion', 'EncryptedProfileBundle':'AES-256-GCM profile storage and platform crypto adapters',
 'EyeCropExtractor':'extracts normalized eye ROI crops for audit/annotation', 'EyeLandmarks':'named eye landmark structures', 'FaceBackend':'common detector/landmark backend contract',
 'FaceDetector':'OpenCV YuNet face detector wrapper', 'FaceMediaPipeRuntime':'runtime loads the optional MediaPipe C ABI',
 'LbfEyeLandmarkMapper':'maps OpenCV LBF 68-point eyes', 'LbfLandmarkDetector':'OpenCV facemark LBF wrapper', 'MediaPipeBackend':'FaceBackend implementation for runtime MediaPipe',
 'MediaPipeEyeLandmarkMapper':'maps MediaPipe topology to semantic eyes', 'PfldEyeLandmarkMapper':'maps PFLD topology to semantic eyes', 'PfldLandmarkProvider':'OpenCV-DNN PFLD inference wrapper',
 'ProcessingRoiAdapter':'crops configured processing ROI and restores coordinates', 'RecordedFrameClock':'monotonic source-frame timestamps', 'ResourceProfiler':'periodic process/core/RSS/thread/phase profiler',
 'YuNetLbfBackend':'real-time YuNet plus LBF backend', 'YuNetPfldBackend':'candidate YuNet plus PFLD backend',
 'main':'interactive camera/video demo and overlay integration', 'benchmark_main':'deterministic benchmark, trace, statistics and resource evidence driver',
 'driver_profile_admin':'offline encrypted profile administration CLI', 'recognition_baseline':'SFace public-fixture recognition baseline CLI', 'pad_baseline':'anti-spoof model diagnostic baseline CLI',
 'sponsor_selftest':'packaged deterministic sponsor self-test executable', 'FaceMediaPipe':'MediaPipe Face Landmarker C ABI implementation',
 'BgrToRgb':'bridge pixel-format conversion', 'LandmarkConversion':'converts MediaPipe landmark results to stable C records', 'MonotonicTimestamp':'strict timestamp monotonicity helper'
}

def describe(path):
    p=path.as_posix(); stem=path.stem
    if stem in PURPOSE: return PURPOSE[stem]
    if p.startswith('tests/'):
        target=stem.removesuffix('_test').replace('_',' ')
        return f"deterministic regression coverage for {target}"
    if p.startswith('scripts/'):
        s=stem.replace('_',' ')
        if 'package' in s: return f"builds a target package and records its payload manifest ({s})"
        if 'verify' in s or 'validate' in s: return f"fail-fast validation workflow ({s})"
        if 'score' in s or 'analyze' in s or 'aggregate' in s: return f"offline evaluation/reporting workflow ({s})"
        if 'mediapipe' in s: return f"MediaPipe acquisition/build/patch workflow ({s})"
        return f"development automation for {s}"
    if p.startswith('mediapipe/'):
        return {'BUILD.bazel':'Bazel bridge build graph','MEDIAPIPE_VERSION':'pinned MediaPipe tag and commit','FaceLandmarkerModel.cmake':'task-model filename and checksum pin'}.get(path.name,'MediaPipe bridge build or test support')
    if p.startswith('config/'): return 'example, documented runtime presentation/ROI/statistics profile'
    if path.name=='CMakeLists.txt': return 'primary C++ build graph, feature flags, executable and CTest registration'
    if path.name=='.bazelversion': return 'pins the Bazel version for the optional MediaPipe bridge'
    if path.name=='README.md': return 'minimal build entry point; use the handoff and architecture docs for full context'
    return 'repository support file'

def guide(tree_img):
    d=Document(); setup(d,"DMS Next Developer Source Guide","Repository map, file responsibilities and safe extension points",True)
    d.add_paragraph("Purpose",style='Heading 1'); d.add_paragraph("A working map for a developer joining the DMS project. It explains the runtime flow, identifies each maintained code/build/test file, and links changes to the validation evidence that must move with them.")
    add_picture(d,tree_img,10.1)
    d.add_heading('Start here',1)
    add_table(d,['Order','Read / run','Why'],[
      ('1','docs/COLLABORATOR_TESTING_HANDOFF.md','Repository rules, branch discipline, claims and private-data boundaries.'),
      ('2','docs/ARCHITECTURE.md and docs/DMS_CORE.md','System boundaries, timestamp/quality invariants and FSM contracts.'),
      ('3','CMakeLists.txt','Targets, optional features, platform dependencies and tests.'),
      ('4','Configure a clean Release build; run ctest','Establish a known baseline before modifying behavior.'),
      ('5','Create a purpose-specific feature branch','Keep Stage 20/21/23/24 evidence and release decisions isolated.')],font_size=8.5)
    d.add_heading('Runtime mental model',1)
    d.add_paragraph("Capture produces source-timestamped frames. A selected backend performs detection/landmarks. Topology adapters convert provider indices to semantic observations. Quality gates decide whether evidence is usable. Calibration, metrics and deterministic temporal FSMs operate on those observations. Rendering, benchmark traces, statistics and future alert adapters consume snapshots; they do not own policy thresholds.")
    d.add_heading('Build products',1)
    add_table(d,['Target','Role','Primary files'],[
      ('yunet_demo','Interactive monitoring/demo','src/main.cpp plus backend, display and dms_core libraries'),
      ('face_benchmark','Recorded/live deterministic evidence generator','src/benchmark_main.cpp, ResourceProfiler, semantic adapters'),
      ('driver_profile_admin','Separate offline profile/enrollment administration','src/driver_profile_admin.cpp, dms_core'),
      ('face_recognition_baseline / face_pad_baseline','Evaluation-only model plumbing','src/recognition_baseline.cpp; src/pad_baseline.cpp'),
      ('dms_sponsor_selftest','Packaged non-camera health check','src/sponsor_selftest.cpp'),
      ('FaceMediaPipe','Optional Bazel-built runtime bridge','mediapipe/api and mediapipe/src')],font_size=8.5)
    files=git_files()
    groups=[('Public interfaces — include/',lambda p:p.as_posix().startswith('include/')),
            ('C++ implementations and executables — src/',lambda p:p.as_posix().startswith('src/')),
            ('Optional MediaPipe bridge — mediapipe/',lambda p:p.as_posix().startswith('mediapipe/')),
            ('Regression and integration tests — tests/',lambda p:p.as_posix().startswith('tests/')),
            ('Automation — scripts/',lambda p:p.as_posix().startswith('scripts/')),
            ('Runtime configuration and build entry points',lambda p:p.as_posix().startswith('config/') or p.name in ['CMakeLists.txt','.bazelversion','README.md'])]
    for title,pred in groups:
        d.add_page_break(); d.add_heading(title,1)
        rows=[]
        for p in files:
            if pred(p):
                kind='Contract' if p.suffix in ['.hpp','.h'] else ('Test' if 'test' in p.as_posix() else 'Implementation / tool')
                note='Change with its paired implementation and focused test.' if p.as_posix().startswith('include/') else ('Private inputs and generated outputs stay outside Git.' if p.as_posix().startswith('scripts/') else 'Keep behavior deterministic and update evidence for changed claims.')
                rows.append((p.as_posix(),kind,describe(p),note))
        add_table(d,['File','Kind','Responsibility','Change discipline'],rows,[2.3,1.0,4.6,2.2],7.4)
    d.add_page_break(); d.add_heading('Documentation and evidence map',1)
    docrows=[]
    for p in files:
        if p.as_posix().startswith('docs/'):
            n=p.name
            if 'STAGE20' in n: area='Stage 20 accuracy/data evidence'
            elif 'STAGE21' in n: area='Stage 21 driver identification'
            elif 'RESOURCE' in n or 'PERFORMANCE' in n: area='Performance/resource evidence'
            elif 'SPONSOR' in n: area='Sponsor demo/feedback'
            elif 'ARCHITECTURE' in n or 'DMS_CORE' in n: area='Architecture/contracts'
            else: area='Handoff, platform or release readiness'
            docrows.append((p.as_posix(),area))
    add_table(d,['Document','Use'],docrows,[5.0,5.1],8)
    d.add_heading('Feature-development checklist',1)
    for s in ["Write the observable contract and failure behavior before choosing a provider.","Keep provider landmark indices behind semantic adapters.","Use source-frame monotonic time; do not derive events from UI or worker latency.","Treat missing/stale/occluded evidence as Unknown or recovering—not as a negative physiological state.","Add focused deterministic tests, then cross-platform build/tests proportional to the claim.","Update the relevant stage report and architecture/inventory when behavior, dependencies or packaging change.","Never commit participant media, identity data, raw private traces, credentials or generated output directories."]:
        d.add_paragraph(s,style='List Number')
    d.add_heading('Current branch and project cautions',1)
    d.add_paragraph("Stage work is intentionally split across feature branches. Stage 24 is the most recent display/ROI/statistics integration line; Stage 23 has a closed Windows/Orin resource-instrumentation line; Stage 21 driver identification remains isolated. Merge, release, model training and production-threshold changes require explicit approval.")
    out=DOCS/'DMS_DEVELOPER_SOURCE_GUIDE.docx'; d.save(out); return out

def sponsor(user,runtime,portable,tech):
    d=Document(); setup(d,"DMS Next Architecture","Sponsor brief — user value, runtime design, portability and technology evidence")
    d.add_paragraph("Executive view",style='Heading 1')
    d.add_paragraph("DMS Next is a provider-neutral driver-monitoring platform. The reusable C++ core converts timestamped visual observations into quality-gated calibration, eye/blink statistics and temporal driver-state decisions. Camera, inference runtime, operating-system loading, security storage and device telemetry remain replaceable target adapters.")
    add_table(d,['Capability','Current checkpoint','Next product gate'],[
      ('Monitoring core','Implemented and deterministically tested','Representative target-camera accuracy and product thresholds'),
      ('Display / processing ROI / statistics','Stage 24 engineering complete','Approve UI wording, rolling window and vehicle-specific ROI'),
      ('CPU/core/memory evidence','Stage 23 Windows + Orin scope complete','Live driver comparison; Ubuntu x64 deferred'),
      ('Driver identification','Offline interfaces, baseline and encrypted profile foundation','Separate biometric/PAD consent, private evaluation and threshold approval'),
      ('TI AM62','Enablement plan only','Board, camera, Processor SDK and target acceptance run')],font_size=8)
    for title,img,caption in [
      ('1. User workflow',user,'Enrollment and identity are gated; Unknown is a valid outcome. Calibration accepts only stable, usable evidence before monitoring and statistics begin.'),
      ('2. Developer/runtime view',runtime,'Provider-specific indices and SDK types stop before semantic observations. Temporal logic and presentation stay deterministic and independently testable.'),
      ('3. Portability boundary',portable,'Platform independence starts at the C++ semantic core. Native camera, loader, cryptography, packaging and telemetry adapters are validated per target.'),
      ('4. Technology and release evidence',tech,'Versions and checksums support reproducibility; model provenance, notices and legal approval remain separate release gates.')]:
        d.add_page_break(); d.add_heading(title,1); add_picture(d,img,6.9); d.add_paragraph(caption)
    d.add_heading('Design principles',2)
    principles={
      '1. User workflow':['Explicit consent for biometric enrollment.','Open-set recognition returns Unknown when evidence is insufficient.','Driver change/long absence starts a new identity, calibration and statistics epoch.'],
      '2. Developer/runtime view':['Source-frame time—not inference or UI delay—drives events.','Depth-one scheduling bounds latency by dropping superseded work.','Occlusion and stale evidence enter Unknown/recovering states.'],
      '3. Portability boundary':['C++17 core and CMake build are shared.','MediaPipe is an optional five-function runtime-loaded C ABI.','Target support is claimed only after native build, camera, resource and accuracy gates.'],
      '4. Technology and release evidence':['OpenCV 4.8 is the validated baseline; MediaPipe is pinned to v0.10.33/commit.','Model checksums prove byte identity, not redistribution rights.','A release needs exact payload SBOM, notices, model provenance and owner/legal approval.']}
    for k,vals in principles.items():
        d.add_heading(k,3)
        for v in vals:d.add_paragraph(v,style='List Bullet')
    d.add_page_break(); d.add_heading('Platform and technology summary',1)
    add_table(d,['Layer / component','Version / boundary','Platforms','License / release note'],[
      ('DMS application core','C++17; provider-neutral semantics','Windows x64, Ubuntu x64*, Orin aarch64','Project license not yet declared; release blocked pending owner selection.'),
      ('OpenCV + contrib','Validated baseline 4.8.0','Windows, Linux x64/aarch64','Apache-2.0 plus exact-build third-party notices.'),
      ('MediaPipe bridge','0.10.33; commit 3987048…; C ABI v1','Packaged on Windows x64 and Orin','Apache-2.0 plus transitive dependency notices.'),
      ('YuNet / LBF / Face Landmarker models','Repository checksum-pinned','Per provider/target evidence','Exact artifact provenance and redistribution terms remain release blockers.'),
      ('Encrypted driver profiles','AES-256-GCM; PBKDF2-HMAC-SHA-256','Windows CNG/DPAPI; Linux OpenSSL 3','Identity processing remains separately consented and gated.'),
      ('TI SK-AM62','Planned adapter/package','Not yet validated','Select BSP/SDK and inventory TI/OSS terms before claim.')],font_size=7.7)
    d.add_paragraph('* Generic Ubuntu x64 was validated earlier for applicable core/provider work; the Stage 23 collector re-check is explicitly deferred because the available WSL image lacks a compiler.')
    d.add_heading('Sponsor decisions requested',1)
    for s in ['Approve final display wording, five-minute rolling-window default and reset behavior at confirmed driver change.','Provide a representative driver/camera session for live Windows and Orin measurements.','For identification, approve separate biometric/PAD consent and later production thresholds after private results.','For AM62, provide the physical SK-AM62, intended camera and selected Processor SDK.','Select the project source license and authorize legal/release review when distribution is intended.']:
        d.add_paragraph(s,style='List Number')
    d.add_heading('Claim boundary',1); d.add_paragraph('This brief describes committed engineering evidence, not a safety-certified product, release authorization or production-accuracy claim. Merge, release, model training and threshold changes remain separately controlled.')
    out=DOCS/'DMS_SPONSOR_ARCHITECTURE_BRIEF.docx'; d.save(out); return out

if __name__=='__main__':
    imgs=diagrams(); outputs=[guide(imgs[4]),sponsor(*imgs[:4])]
    print('\n'.join(str(x) for x in outputs))
